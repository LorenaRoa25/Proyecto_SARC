"""Inserta mockups oficiales anotados en el Anexo B del documento IEEE."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw


PROJECT = Path(__file__).resolve().parents[1]
DOC_PATH = PROJECT / "Documentos" / "SARC_SRS_IEEE830_ICONTEC.docx"
ZIP_PATH = Path(r"C:\Users\Usuario\Downloads\ilovepdf_pages-to-jpg.zip")
EXTRACT_DIR = PROJECT / "Documentos" / "mockups_extraidos"
ANNOTATED_DIR = PROJECT / "Documentos" / "mockups_anotados"

BLUE = (11, 79, 138)
GREEN = (106, 168, 79)
RED = (207, 42, 39)
YELLOW = (255, 197, 61)


FIGURES = [
    {
        "key": "Login",
        "page": "0001",
        "caption": "Figura 1. Pantalla de inicio de sesión del sistema SARC.",
        "desc": "Usar la imagen correspondiente al mockup de inicio de sesión proporcionado por el usuario.",
        "resaltar": "Campo de correo, campo de contraseña, checkbox de Habeas Data y botón INGRESAR.",
        "anotar": "Recuadros en campos y botón; flecha suave indicando el flujo correo -> contraseña -> aceptación -> ingreso.",
        "boxes": [(610, 520, 1145, 580), (610, 600, 1145, 660), (600, 675, 1155, 735), (610, 745, 1145, 815)],
        "arrows": [((880, 585), (880, 600)), ((880, 665), (880, 680)), ((880, 735), (880, 750))],
    },
    {
        "key": "Dashboard",
        "page": "0006",
        "caption": "Figura 2. Dashboard principal del sistema SARC.",
        "desc": "Usar el mockup del panel principal con bienvenida, cursos sugeridos y asistente virtual.",
        "resaltar": "Banner de bienvenida, menú superior, lista de cursos sugeridos y panel del asistente.",
        "anotar": "Recuadros en las zonas principales y flecha desde el menú hacia los módulos.",
        "boxes": [(30, 140, 1680, 250), (35, 260, 860, 885), (880, 260, 1680, 885), (300, 95, 1450, 140)],
        "arrows": [((620, 130), (520, 265)), ((1120, 130), (1230, 265))],
    },
    {
        "key": "Catalogo",
        "page": "0008",
        "caption": "Figura 3. Catálogo de cursos sugeridos.",
        "desc": "Usar el mockup del módulo Recomendaciones con filtros y listado de cursos.",
        "resaltar": "Filtros de área/modalidad, filas de cursos, botones Ver detalles e INSCRIBIRSE.",
        "anotar": "Recuadros en filtros y botones; flecha hacia el flujo de consulta e inscripción.",
        "boxes": [(40, 205, 1680, 285), (60, 350, 1660, 520), (60, 535, 1660, 705), (60, 720, 1660, 890)],
        "arrows": [((1180, 400), (1380, 400)), ((1180, 585), (1380, 585)), ((1180, 770), (1380, 770))],
    },
    {
        "key": "Detalle del curso",
        "page": "0026",
        "caption": "Figura 4. Detalle de curso seleccionado.",
        "desc": "Usar el mockup de detalle del curso Matemáticas Básicas.",
        "resaltar": "Nombre del curso, descripción, duración, nivel, botón VOLVER y botón INSCRIBIRSE.",
        "anotar": "Recuadros en información académica y acciones principales.",
        "boxes": [(35, 145, 1680, 230), (90, 310, 1550, 610), (250, 770, 520, 850), (1375, 770, 1630, 850)],
        "arrows": [((760, 615), (445, 770)), ((1060, 615), (1500, 770))],
    },
    {
        "key": "Progreso académico",
        "page": "0038",
        "caption": "Figura 5. Vista de progreso académico.",
        "desc": "Usar el mockup de Mi progreso con cursos, barras y botón de reporte.",
        "resaltar": "Barras de progreso, estados de cursos y botones Ver reporte detallado.",
        "anotar": "Recuadros en las filas de avance y círculos en botones de reporte.",
        "boxes": [(40, 250, 1680, 430), (40, 455, 1680, 635), (40, 660, 1680, 850)],
        "circles": [(1470, 335, 120), (1470, 540, 120), (1470, 745, 120)],
    },
    {
        "key": "Reporte PDF",
        "page": "0040",
        "caption": "Figura 6. Reporte PDF de progreso académico.",
        "desc": "Usar el mockup del reporte PDF generado por el sistema.",
        "resaltar": "Título del reporte, datos del estudiante, progreso, promedio y recomendación.",
        "anotar": "Recuadro general sobre el documento PDF y resaltado suave en la información académica.",
        "boxes": [(520, 180, 1330, 1000), (590, 350, 1250, 820)],
        "arrows": [((450, 520), (590, 520))],
    },
    {
        "key": "Perfil",
        "page": "0045",
        "caption": "Figura 7. Perfil de usuario.",
        "desc": "Usar el mockup del perfil con datos personales y acciones de cuenta.",
        "resaltar": "Datos del estudiante, avatar, botón Editar Perfil y botón Cerrar Sesión.",
        "anotar": "Recuadros en datos y botones; círculo en avatar.",
        "boxes": [(80, 300, 850, 660), (260, 780, 520, 860), (1290, 780, 1580, 860)],
        "circles": [(1330, 430, 140)],
    },
    {
        "key": "Recuperacion",
        "page": "0004",
        "caption": "Figura 8. Recuperación de contraseña.",
        "desc": "Usar el mockup de recuperación de contraseña.",
        "resaltar": "Campo de correo, campos de contraseña y botones Volver/Enviar.",
        "anotar": "Recuadros en campos y botones; flecha hacia el envío de recuperación.",
        "boxes": [(650, 455, 1120, 525), (650, 545, 1120, 615), (650, 635, 1120, 705), (650, 725, 1120, 810)],
        "arrows": [((885, 705), (885, 735))],
    },
    {
        "key": "Recomendaciones de cursos",
        "page": "0020",
        "caption": "Figura 9. Recomendaciones de cursos y filtros.",
        "desc": "Usar el mockup del catálogo con selector de área interdisciplinaria desplegado.",
        "resaltar": "Filtro de área, filtro de modalidad y cursos sugeridos.",
        "anotar": "Recuadro en el desplegable y resaltado en la zona de resultados.",
        "boxes": [(50, 200, 470, 505), (1250, 200, 1650, 275), (55, 350, 1660, 890)],
        "arrows": [((470, 300), (760, 390))],
    },
    {
        "key": "Asistente virtual",
        "page": "0007",
        "caption": "Figura 10. Asistente virtual de apoyo académico.",
        "desc": "Usar el mockup del dashboard donde aparece el panel Asistente Virtual.",
        "resaltar": "Título Asistente Virtual y lista de recomendaciones académicas.",
        "anotar": "Recuadro en el panel del asistente; círculo discreto sobre las recomendaciones.",
        "boxes": [(880, 260, 1680, 885)],
        "circles": [(1180, 520, 230)],
        "arrows": [((830, 520), (895, 520))],
    },
]


def extract_zip() -> None:
    EXTRACT_DIR.mkdir(exist_ok=True)
    if not any(EXTRACT_DIR.glob("*.jpg")):
        with ZipFile(ZIP_PATH) as zf:
            zf.extractall(EXTRACT_DIR)


def image_for_page(page: str) -> Path:
    matches = sorted(EXTRACT_DIR.glob(f"*page-{page}.jpg"))
    if not matches:
        raise FileNotFoundError(f"No se encontro page-{page}")
    return matches[0]


def arrow(draw: ImageDraw.ImageDraw, start, end, fill, width=8):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    left = (x2 - 28 * ux - 14 * uy, y2 - 28 * uy + 14 * ux)
    right = (x2 - 28 * ux + 14 * uy, y2 - 28 * uy - 14 * ux)
    draw.polygon([end, left, right], fill=fill)


def annotate_images() -> None:
    ANNOTATED_DIR.mkdir(exist_ok=True)
    for item in FIGURES:
        src = image_for_page(item["page"])
        im = Image.open(src).convert("RGB")
        draw = ImageDraw.Draw(im, "RGBA")
        for box in item.get("boxes", []):
            draw.rectangle(box, outline=RED + (255,), width=8)
            x1, y1, x2, y2 = box
            draw.rectangle((x1, y1, x2, y2), fill=YELLOW + (24,))
        for center_x, center_y, radius in item.get("circles", []):
            draw.ellipse((center_x - radius, center_y - radius, center_x + radius, center_y + radius), outline=BLUE + (255,), width=8)
        for start, end in item.get("arrows", []):
            arrow(draw, start, end, GREEN + (255,), width=8)
        out = ANNOTATED_DIR / f"{item['key'].lower().replace(' ', '_')}.jpg"
        im.save(out, quality=92)
        item["annotated"] = out


def paragraph_text(element):
    return "".join(node.text or "" for node in element.iter() if node.tag.endswith("}t"))


def find_heading(doc: Document, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading") and paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(prefix)


def child_index(body, element):
    children = list(body.iterchildren())
    for index, child in enumerate(children):
        if child is element:
            return index
    raise ValueError("elemento no encontrado")


def add_run(paragraph, text, bold=False):
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(31, 41, 55)
    run.bold = bold
    return run


def add_annex_content(doc: Document):
    doc.add_heading("Anexo B. Análisis de mockups y prototipos", level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_run(p, "Las capturas de este anexo provienen exclusivamente del archivo de mockups/prototipos suministrado por el usuario. No se generaron capturas nuevas ni screenshots reales de la aplicación en ejecución.")

    doc.add_heading("B.1 Capturas requeridas para mockups", level=3)
    for index, item in enumerate(FIGURES, start=1):
        marker = doc.add_paragraph()
        marker.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_run(marker, f"[INSERTAR MOCKUP — {item['key']}]", bold=True)

        for label, key in [
            ("Imagen oficial", "desc"),
            ("Elementos a resaltar", "resaltar"),
            ("Anotaciones visuales", "anotar"),
        ]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_run(p, f"{label}: ", bold=True)
            add_run(p, item[key])

        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = pic.add_run()
        run.add_picture(str(item["annotated"]), width=Inches(5.75))

        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = caption.add_run(item["caption"])
        r.italic = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        if index < len(FIGURES):
            doc.add_paragraph()


def replace_annex_b() -> None:
    doc = Document(DOC_PATH)
    body = doc.element.body

    start_p = find_heading(doc, "Anexo B.")
    end_p = find_heading(doc, "Anexo C.")
    children = list(body.iterchildren())
    start = child_index(body, start_p._p)
    end = child_index(body, end_p._p)

    for element in children[start:end]:
        body.remove(element)

    children = list(body.iterchildren())
    insert_at = child_index(body, end_p._p)
    len_before = len(children)
    add_annex_content(doc)
    children_after = list(body.iterchildren())
    new_elements = children_after[len_before - 1 : -1]
    for element in new_elements:
        body.remove(element)
    for offset, element in enumerate(new_elements):
        body.insert(insert_at + offset, element)

    doc.save(DOC_PATH)


def main():
    extract_zip()
    annotate_images()
    replace_annex_b()
    print(DOC_PATH)


if __name__ == "__main__":
    main()
