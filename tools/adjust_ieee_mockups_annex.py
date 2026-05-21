"""Ajusta el documento IEEE: mueve mockups a anexos y corrige numeracion."""

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOC_PATH = r"Documentos\SARC_SRS_IEEE830_ICONTEC.docx"


def is_paragraph(element):
    return element.tag.endswith("}p")


def paragraph_text(element):
    return "".join(node.text or "" for node in element.iter() if node.tag.endswith("}t"))


def paragraph_style(element):
    p_pr = element.find(qn("w:pPr"))
    if p_pr is None:
        return ""
    p_style = p_pr.find(qn("w:pStyle"))
    if p_style is None:
        return ""
    return p_style.get(qn("w:val")) or ""


def set_text(paragraph: Paragraph, text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def is_heading_style(style: str, level: int | None = None):
    if not style:
        return False
    if level is None:
        return style.startswith("Heading") or style.startswith("Ttulo")
    return style in {f"Heading{level}", f"Heading {level}", f"Ttulo{level}"}


def find_child_index(children, startswith: str, heading_level: int | None = None):
    for index, child in enumerate(children):
        if not is_paragraph(child):
            continue
        text_ok = paragraph_text(child).strip().startswith(startswith)
        style_ok = True if heading_level is None else is_heading_style(paragraph_style(child), heading_level)
        if text_ok and style_ok:
            return index
    raise ValueError(f"No se encontro: {startswith}")


def paragraph_by_prefix(doc: Document, prefix: str, heading_only: bool = False):
    for paragraph in doc.paragraphs:
        if heading_only and not paragraph.style.name.startswith("Heading"):
            continue
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"No se encontro parrafo: {prefix}")


def main():
    doc = Document(DOC_PATH)
    body = doc.element.body
    children = list(body.iterchildren())

    mockups_start = find_child_index(children, "7. An", 1)
    annex_start = find_child_index(children, "10. Anexos", 1)

    mockup_block = children[mockups_start:annex_start]
    for element in mockup_block:
        body.remove(element)

    # Reinsertar el bloque antes de las historias de usuario, ahora como Anexo B.
    children = list(body.iterchildren())
    insert_before = find_child_index(children, "Anexo B.", 2)
    for offset, element in enumerate(mockup_block):
        body.insert(insert_before + offset, element)

    # Ajustar textos principales.
    p = paragraph_by_prefix(doc, "10. Anexos", heading_only=True)
    set_text(p, "7. Anexos")
    p.style = doc.styles["Heading 1"]

    p = paragraph_by_prefix(doc, "7. An", heading_only=True)
    set_text(p, "Anexo B. Análisis de mockups y prototipos")
    p.style = doc.styles["Heading 2"]

    p = paragraph_by_prefix(doc, "7.1 Capturas", heading_only=True)
    set_text(p, "B.1 Capturas requeridas para mockups")
    p.style = doc.styles["Heading 3"]

    # Correr anexos posteriores.
    replacements = {
        "Anexo B. Historias": "Anexo C. Historias de usuario SCRUM",
        "Anexo C. Ficha": "Anexo D. Ficha técnica del software y hardware",
        "Anexo D. Cronograma": "Anexo E. Cronograma de Gantt",
    }
    for prefix, new_text in replacements.items():
        p = paragraph_by_prefix(doc, prefix, heading_only=True)
        set_text(p, new_text)
        p.style = doc.styles["Heading 2"]

    doc.save(DOC_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
