"""Extrae texto de la rúbrica PDF y documentos DOCX para el análisis."""

from pathlib import Path

from docx import Document
from pypdf import PdfReader


PROJECT_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_DIR / "analysis_inputs"

FILES = {
    "escala": Path(r"C:\Users\Usuario\Downloads\ESCALA DE EVALUACIÓN PROYECTOS UNISABANETA.pdf"),
    "proyecto": Path(r"C:\Users\Usuario\OneDrive\Desktop\Documento proyecto SARC - Lorena Roa Rivera.docx"),
    "minuta": Path(r"C:\Users\Usuario\OneDrive\Desktop\Minuta Licenciamiento-Lorena Roa.docx"),
}


def extract_pdf(path):
    """Extrae texto por página desde un PDF."""
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(f"\n--- PAGE {index} ---\n{page.extract_text() or ''}")
    return "\n".join(pages), len(reader.pages)


def extract_docx(path):
    """Extrae párrafos y tablas desde un documento Word."""
    document = Document(str(path))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"\n--- TABLE {table_index} ---")
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            parts.append(" | ".join(values))

    return "\n".join(parts)


def main():
    """Genera archivos TXT auxiliares para lectura y comparación."""
    OUTPUT_DIR.mkdir(exist_ok=True)

    pdf_text, page_count = extract_pdf(FILES["escala"])
    (OUTPUT_DIR / "escala_evaluacion.txt").write_text(pdf_text, encoding="utf-8")
    (OUTPUT_DIR / "documento_proyecto_sarc.txt").write_text(extract_docx(FILES["proyecto"]), encoding="utf-8")
    (OUTPUT_DIR / "minuta_licenciamiento.txt").write_text(extract_docx(FILES["minuta"]), encoding="utf-8")

    print(f"Textos extraídos en: {OUTPUT_DIR}")
    print(f"Páginas PDF: {page_count}")


if __name__ == "__main__":
    main()
