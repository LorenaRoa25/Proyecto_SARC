"""Corrige encabezados finales de anexos en el documento IEEE."""

from docx import Document


DOC_PATH = r"Documentos\SARC_SRS_IEEE830_ICONTEC.docx"


def set_text(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def main():
    doc = Document(DOC_PATH)

    fixed_annex_heading = False
    fixed_mockup_heading = False

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if (
            not fixed_annex_heading
            and paragraph.style.name.startswith("Heading")
            and text.startswith("Anexo B. An")
        ):
            set_text(paragraph, "7. Anexos")
            paragraph.style = doc.styles["Heading 1"]
            fixed_annex_heading = True
            continue

        if (
            not fixed_mockup_heading
            and paragraph.style.name.startswith("Heading")
            and text.startswith("7. An")
        ):
            set_text(paragraph, "Anexo B. Análisis de mockups y prototipos")
            paragraph.style = doc.styles["Heading 2"]
            fixed_mockup_heading = True

    doc.save(DOC_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
