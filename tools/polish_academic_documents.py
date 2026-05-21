"""Aplica ajustes academicos finales a documentos SARC existentes."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_DIR / "Documentos"
PROJECT_DOC = DOCS_DIR / "Documento proyecto SARC - Lorena Roa Rivera.docx"
REVISION_MD = DOCS_DIR / "REVISION_DOCUMENTAL_SARC.md"

BLUE = "0B4F8A"
GREEN = "6AA84F"
LIGHT_BLUE = "EAF3FA"
TEXT = RGBColor(31, 41, 55)
GITHUB_URL = "https://github.com/LorenaRoa25/Proyecto_SARC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = TEXT
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_existing_tables(doc: Document) -> None:
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        if run.font.size is None:
                            run.font.size = Pt(8)
                if row_index == 0:
                    set_cell_shading(cell, BLUE)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.bold = True


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor.from_string(BLUE)


def add_para(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.color.rgb = TEXT


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=8)
        set_cell_shading(table.rows[0].cells[i], BLUE)
        table.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            table.rows[0].cells[i].width = Inches(widths[i])
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            set_cell_text(row.cells[i], value, size=8)
            if i == 0:
                set_cell_shading(row.cells[i], LIGHT_BLUE)
            if widths:
                row.cells[i].width = Inches(widths[i])
    doc.add_paragraph()


def replace_all_text(doc: Document, replacements: dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                if paragraph.runs:
                    paragraph.runs[0].text = paragraph.text.replace(old, new)
                    for run in paragraph.runs[1:]:
                        run.text = ""
                else:
                    paragraph.add_run(paragraph.text.replace(old, new))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for old, new in replacements.items():
                        if old in paragraph.text:
                            if paragraph.runs:
                                paragraph.runs[0].text = paragraph.text.replace(old, new)
                                for run in paragraph.runs[1:]:
                                    run.text = ""
                            else:
                                paragraph.add_run(paragraph.text.replace(old, new))


def project_doc_has_github_annex(doc: Document) -> bool:
    return any(GITHUB_URL in paragraph.text for paragraph in doc.paragraphs) or any(
        GITHUB_URL in cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )


def update_project_document() -> None:
    doc = Document(PROJECT_DOC)
    replace_all_text(doc, {
        "credenciales institucionales": "credenciales gestionadas mediante Firebase Authentication",
        "registro e inicio de sesion": "inicio de sesion",
        "registro e inicio de sesión": "inicio de sesión",
        "Registro e inicio de sesion": "Inicio de sesion",
        "Registro e inicio de sesión": "Inicio de sesión",
        "administracion de usuarios": "gestion de usuarios desde Firebase y perfiles futuros",
        "administración de usuarios": "gestión de usuarios desde Firebase y perfiles futuros",
    })
    style_existing_tables(doc)

    if not project_doc_has_github_annex(doc):
        doc.add_page_break()
        add_heading(doc, "12. ANEXO - REPOSITORIO GITHUB", 1)
        add_para(
            doc,
            "Este anexo documenta el repositorio oficial del proyecto SARC como evidencia tecnica de control de versiones, trazabilidad del desarrollo y organizacion del codigo fuente. No se limita a un enlace: describe su objetivo, contenido y valor academico dentro del proceso de ingenieria de software.",
        )
        add_table(doc, ["Campo", "Descripcion"], [
            ["Nombre del repositorio", "Proyecto_SARC"],
            ["Descripcion", "Repositorio del Sistema de Apoyo Academico con Recomendacion de Cursos, aplicacion web modular orientada al apoyo academico universitario."],
            ["Objetivo", "Centralizar codigo fuente, configuracion Firebase, reglas de Firestore, documentos academicos y herramientas de generacion documental."],
            ["Enlace oficial", GITHUB_URL],
            ["Contenido principal", "index.html, app.py, pages/, shared/, firestore.rules, README.md, FIREBASE_SETUP.md, Documentos/ y tools/."],
            ["Importancia", "Permite mantener historial de cambios, recuperar versiones, sustentar decisiones tecnicas y facilitar futuras revisiones o despliegues."],
        ], [1.7, 5.1])

        add_heading(doc, "12.1 Coherencia funcional documentada", 2)
        add_table(doc, ["Aspecto revisado", "Estado actualizado"], [
            ["Inicio de sesion", "El acceso exitoso entra al sistema sin mostrar una notificacion verde automatica de sesion autenticada."],
            ["Perfil demo", "El boton Reiniciar demo solo corresponde al usuario lorena.roa@unisabaneta.edu.co."],
            ["Usuarios nuevos", "No deben visualizar funciones exclusivas de demostracion; la validacion depende del correo autorizado, no solo de tener sesion."],
            ["Firebase", "Authentication gestiona credenciales; Firestore almacena datos academicos por userId."],
        ], [2.0, 4.8])

    core = doc.core_properties
    core.title = "Documento principal SARC"
    core.subject = "Sistema de Apoyo Academico con Recomendacion de Cursos"
    core.keywords = "SARC, Firebase, GitHub, SRS, Gantt, historias de usuario"
    doc.save(PROJECT_DOC)


def update_revision_markdown() -> None:
    content = f"""# Revision documental SARC

Fecha de revision: {date.today().isoformat()}

## Alcance revisado

Se revisaron los documentos principales ubicados en la carpeta `Documentos`:

- `DOCUMENTACION_CODIGO_SARC.docx`
- `Documento proyecto SARC - Lorena Roa Rivera.docx`
- `Guia_Rapida_Uso_SARC.docx`
- `Informe_evaluacion_SARC.docx`
- `Manual_Tecnico_SARC.docx`
- `Manual_Usuario_SARC.docx`
- `Minuta Licenciamiento-Lorena Roa.docx`
- `Resumen_Ejecutivo_Minuta_SARC.docx`
- `SARC_SRS_IEEE830_ICONTEC.docx`

Tambien se verifico la existencia de carpetas de capturas y mockups usados como soporte visual academico.

## Ajustes aplicados

- Cronograma de Gantt reorganizado como matriz visual por semanas, con colores institucionales de SARC.
- Historias de usuario mejoradas en formato SCRUM academico con ID, historia, criterios, prioridad y estimacion.
- Ficha tecnica de hardware separada de la ficha tecnica de software.
- Anexo de Repositorio GitHub agregado al documento principal y al SRS.
- Manual Tecnico reforzado con tecnologias, estructura de carpetas, Firebase, flujo del sistema, autenticacion, almacenamiento y perfiles demo.
- Documentacion alineada con el cambio reciente de login limpio sin notificacion verde de exito.
- Funcion Reiniciar demo documentada como exclusiva de `lorena.roa@unisabaneta.edu.co`.

## Estado funcional documentado

SARC corresponde a una aplicacion web academica modular desarrollada con HTML5, CSS3 y JavaScript ES Modules, con autenticacion mediante Firebase Authentication y persistencia en Cloud Firestore. El sistema incluye login, recuperacion, dashboard, recomendaciones, catalogo filtrable, detalle de curso, inscripcion, progreso, reportes PDF y perfil.

## Criterio de coherencia

La documentacion diferencia lo implementado de lo proyectado: no existe panel administrativo, no hay chatbot conversacional ni machine learning avanzado, y el registro formal de usuarios queda como evolucion futura. La linea visual se mantiene con paleta azul y verde institucional, tablas limpias y estilo universitario profesional.
"""
    REVISION_MD.write_text(content, encoding="utf-8")


def main() -> None:
    update_project_document()
    update_revision_markdown()
    print("Documentos academicos pulidos.")


if __name__ == "__main__":
    main()
