"""Genera una guia de sustentacion del codigo SARC en formato DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_DIR / "GUIA_SUSTENTACION_CODIGO_SARC.docx"

BLUE = RGBColor(8, 83, 148)
GREEN = RGBColor(106, 168, 79)
GRAY = RGBColor(90, 90, 90)
LIGHT_BLUE = "EAF3FB"
LIGHT_GREEN = "EEF7EA"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Segoe UI"
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margin(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Segoe UI"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08

    for style_name, size in [("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)]:
        style = styles[style_name]
        style.font.name = "Segoe UI"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    styles["List Bullet"].font.name = "Segoe UI"
    styles["List Number"].font.name = "Segoe UI"


def add_title(document):
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("Guia de sustentacion del codigo SARC")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = BLUE

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run(
        "Explicacion tecnica, natural y paso a paso para defender el proyecto ante el profesor"
    )
    run.font.size = Pt(11.5)
    run.font.color.rgb = GRAY


def add_heading(document, text, level=1):
    return document.add_heading(text, level=level)


def add_note(document, title, body, fill=LIGHT_BLUE):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margin(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    title_run = paragraph.add_run(title)
    title_run.bold = True
    title_run.font.color.rgb = BLUE
    title_run.font.name = "Segoe UI"
    title_run.font.size = Pt(10)
    paragraph.add_run("\n")
    body_run = paragraph.add_run(body)
    body_run.font.name = "Segoe UI"
    body_run.font.size = Pt(9.5)
    document.add_paragraph()


def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.4)
    run.font.color.rgb = RGBColor(45, 45, 45)
    return paragraph


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "085394")
        set_cell_margin(cell)
        set_cell_text(cell, header, bold=True, color=RGBColor(255, 255, 255))
        if widths:
            cell.width = Inches(widths[index])

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_margin(cells[index])
            set_cell_text(cells[index], value)
            if widths:
                cells[index].width = Inches(widths[index])

    document.add_paragraph()
    return table


def bullet(document, text):
    document.add_paragraph(text, style="List Bullet")


def number(document, text):
    document.add_paragraph(text, style="List Number")


def polish_spanish_accents(document):
    """Ajusta tildes frecuentes sin cambiar la estructura del documento."""
    replacements = {
        "Guia": "Guía",
        "sustentacion": "sustentación",
        "codigo": "código",
        "Explicacion": "Explicación",
        "explicacion": "explicación",
        "tecnica": "técnica",
        "tecnico": "técnico",
        "rapida": "rápida",
        "aplicacion": "aplicación",
        "academico": "académico",
        "autenticacion": "autenticación",
        "informacion": "información",
        "perdio": "perdió",
        "perderse": "perderse",
        "esta separado": "está separado",
        "Esta construida": "Está construida",
        "Esta separacion": "Esta separación",
        "raiz": "raíz",
        "logica": "lógica",
        "funcion": "función",
        "funciónes": "funciones",
        "funciones": "funciones",
        "navegacion": "navegación",
        "segun": "según",
        "Como": "Cómo",
        "ejecucion": "ejecución",
        "envia": "envía",
        "ingreso": "ingresó",
        "sesion": "sesión",
        "contrasena": "contraseña",
        "contrasenas": "contraseñas",
        "modulos": "módulos",
        "modulo": "módulo",
        "modular": "modular",
        "despues": "después",
        "Asi": "Así",
        "asi": "así",
        "recuperacion": "recuperación",
        "pagina": "página",
        "unico": "único",
        "especifico": "específico",
        "especifica": "específica",
        "practicas": "prácticas",
        "tipicas": "típicas",
        "catalogo": "catálogo",
        "dinamico": "dinámico",
        "dinamicos": "dinámicos",
        "servidor local": "servidor local",
        "autenticado": "autenticado",
        "autenticada": "autenticada",
        "esta autenticado": "está autenticado",
        "esta abierta": "está abierta",
        "esta viendo": "está viendo",
        "esta en": "está en",
        "mas": "más",
        "Tambien": "También",
        "tambien": "también",
        "operacion": "operación",
        "produccion": "producción",
        "vision": "visión",
        "validacion": "validación",
        "Contrase": "Contrase",
        "programacion": "programación",
        "Programacion": "Programación",
        "Ingles": "Inglés",
        "ingles": "inglés",
        "basica": "básica",
        "basico": "básico",
        "metricas": "métricas",
        "linea": "línea",
        "lineas": "líneas",
        "imagenes": "imágenes",
        "usuario": "usuario",
        "multiples": "múltiples",
        "ultima": "última",
        "ultimo": "último",
        "publicas": "públicas",
        "privadas": "privadas",
        "estaticos": "estáticos",
        "estaticas": "estáticas",
        "maqueta": "maqueta",
        "frase": "frase",
        "accion": "acción",
        "acciones": "acciones",
        "notificacion": "notificación",
        "notificaciones": "notificaciones",
        "automaticamente": "automáticamente",
        "dependencias": "dependencias",
        "librerias": "librerías",
    }

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            for old, new in replacements.items():
                run.text = run.text.replace(old, new)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        for old, new in replacements.items():
                            run.text = run.text.replace(old, new)


def build_document():
    doc = Document()
    style_document(doc)
    add_title(doc)

    add_heading(doc, "1. Explicacion rapida para abrir la defensa")
    doc.add_paragraph(
        "SARC es una aplicacion web de apoyo academico que recomienda cursos, permite inscribirse, "
        "muestra el progreso del estudiante y administra datos de perfil. Esta construida con HTML, "
        "CSS y JavaScript modular. La autenticacion se realiza con Firebase Authentication y la "
        "informacion del usuario se guarda en Cloud Firestore."
    )
    add_note(
        doc,
        "Respuesta oral corta",
        "Yo lo explicaria asi: el navegador carga una sola pagina principal, JavaScript decide que vista mostrar "
        "segun la ruta de la URL, cada vista carga su propia plantilla HTML, y los cambios del usuario se guardan "
        "en Firestore asociados a su userId. Asi el proyecto no depende de datos quemados en el navegador.",
        LIGHT_GREEN,
    )

    add_heading(doc, "2. Como leer el proyecto sin perderse")
    doc.add_paragraph(
        "El proyecto esta separado por responsabilidades. La raiz contiene el punto de entrada y el servidor local. "
        "La carpeta shared contiene codigo reutilizable. La carpeta pages contiene cada pantalla. Esta separacion "
        "ayuda a defender el codigo porque se puede decir que cada archivo tiene una tarea concreta."
    )
    add_table(
        doc,
        ["Zona", "Idea clave para explicar"],
        [
            ["index.html", "Es el contenedor principal: carga estilos, crea los espacios de la app y llama app.js."],
            ["shared/js", "Guarda la logica comun: estado, Firebase, rutas, cursos, PDF y plantillas."],
            ["shared/css", "Define estilos base, layout general y componentes reutilizables."],
            ["pages", "Cada vista tiene HTML para estructura, CSS para apariencia y JS para comportamiento."],
            ["firestore.rules", "Protege la base de datos para que cada usuario acceda solo a sus documentos."],
        ],
        widths=[1.6, 4.7],
    )

    add_heading(doc, "3. Flujo general de ejecucion")
    number(doc, "El usuario ejecuta iniciar_sarc.vbs; app.py inicia en segundo plano y el navegador abre http://localhost:8000/.")
    number(doc, "index.html carga todos los CSS y despues importa shared/js/app.js con type=\"module\".")
    number(doc, "app.js ejecuta init(), que inicializa datos, Firebase, sesion, eventos y renderizado.")
    number(doc, "La ruta hash de la URL decide la pantalla: #login, #inicio, #recomendaciones, #detalle/id, #progreso o #perfil.")
    number(doc, "Si el usuario no esta autenticado, app.js solo deja entrar a login o recuperar contrasena.")
    number(doc, "Cuando una vista necesita datos, lee state.db. Cuando modifica algo importante, llama saveDatabase().")
    number(doc, "saveDatabase() envia la copia actual del estado a Firestore mediante firebase-service.js.")

    add_note(
        doc,
        "Frase para sustentar el flujo",
        "El frontend no consulta Firestore desde todas partes. Las vistas trabajan con state.db, y data.js funciona como puente "
        "entre la interfaz y Firebase. Eso reduce repeticion y hace que el codigo sea mas ordenado.",
    )

    add_heading(doc, "4. index.html: el esqueleto de la aplicacion")
    doc.add_paragraph(
        "index.html no contiene la logica del sistema. Su trabajo es preparar el escenario donde JavaScript va a pintar las vistas. "
        "Por eso se le puede llamar el shell principal."
    )
    add_code(
        doc,
        """
<section id="authView" class="auth-shell"></section>
<section id="dashboardView" class="app-shell hidden">
  ...
  <div id="viewContainer" class="view-content"></div>
</section>
<script type="module" src="shared/js/app.js"></script>
""",
    )
    doc.add_paragraph(
        "authView se usa para login y recuperacion. dashboardView se usa cuando el usuario ya ingreso. "
        "viewContainer es el espacio donde se inserta la vista activa. El atributo type=\"module\" permite usar import y export en JavaScript."
    )

    add_heading(doc, "5. app.js: punto de entrada y enrutador")
    doc.add_paragraph(
        "app.js es uno de los archivos mas importantes porque conecta todo: importa funciones de estado, navegacion, componentes y vistas. "
        "Tambien decide que se muestra en pantalla segun la sesion y la ruta."
    )
    add_code(
        doc,
        """
import { enableAllCourses, initializeState, resetStartupSession, state } from "./data.js";
import { navigate } from "./navigation.js";
import { renderLogin } from "../../pages/login/login.js";
import { renderHome } from "../../pages/inicio/inicio.js";
""",
    )
    doc.add_paragraph(
        "Los imports traen funciones desde otros modulos. La idea es no escribir todo en un solo archivo. "
        "app.js no sabe como guardar en Firestore ni como dibujar cada pantalla por dentro; solamente coordina."
    )
    add_code(
        doc,
        """
async function init() {
  await initializeState();
  await resetStartupSession();
  await enableAllCourses();
  bindGlobalEvents();
  setDashboardRefresh(renderDashboard);
  renderApp();
  setTimeout(() => loader.classList.add("is-hidden"), 700);
  window.addEventListener("hashchange", renderApp);
}
""",
    )
    doc.add_paragraph(
        "init() es asincrona porque debe esperar operaciones externas. initializeState() prepara Firebase y los datos. "
        "resetStartupSession() fuerza que la demo inicie desde login. enableAllCourses() evita que los cursos queden sin cupos durante pruebas. "
        "bindGlobalEvents() conecta botones generales. setDashboardRefresh(renderDashboard) le permite a courses.js pedir que el dashboard se vuelva a pintar despues de una inscripcion."
    )
    add_code(
        doc,
        """
if (!state.session.authenticated && !["login", "recuperar"].includes(hash)) {
  window.location.hash = "login";
  return;
}
""",
    )
    doc.add_paragraph(
        "Este bloque es una proteccion de navegacion. Si alguien intenta entrar manualmente a #perfil o #progreso sin sesion, "
        "la app lo devuelve a login. No reemplaza las reglas de Firestore, pero mejora el flujo de interfaz."
    )
    add_code(
        doc,
        """
const pageRenderers = {
  inicio: renderHome,
  recomendaciones: renderRecommendations,
  detalle: renderDetail,
  progreso: renderProgress,
  perfil: renderProfile
};
const renderer = pageRenderers[state.route] || renderHome;
await renderer(viewContainer);
""",
    )
    doc.add_paragraph(
        "Aqui se aplica una idea clave: en vez de usar muchos if repetidos, se crea un objeto que relaciona el nombre de la ruta con su funcion render. "
        "Si la ruta es desconocida, se usa renderHome como respaldo."
    )

    add_heading(doc, "6. data.js: estado global y datos de la aplicacion")
    doc.add_paragraph(
        "data.js centraliza la informacion que usan las vistas. Firestore es la base de datos real, pero state.db es la copia temporal que vive en memoria mientras la app esta abierta."
    )
    add_code(
        doc,
        """
export const state = {
  route: "login",
  currentDetailId: "matematicas",
  filters: { area: "", modality: "" },
  db: null,
  session: null
};
""",
    )
    doc.add_paragraph(
        "state.route guarda la pantalla actual. currentDetailId indica que curso se esta viendo en detalle. "
        "filters guarda los filtros del catalogo. db contiene usuario, cursos y tareas. session indica si hay usuario autenticado."
    )
    add_code(
        doc,
        """
export async function initializeState() {
  await initializeFirebase();
  const user = await waitForAuthUser();
  state.session = loadSession(user);
  state.db = user ? await loadDatabaseForUser(user) : structuredClone(defaultDatabase);
}
""",
    )
    doc.add_paragraph(
        "initializeState() prepara el estado inicial. Primero intenta inicializar Firebase. Luego espera si Firebase ya tiene un usuario activo. "
        "Si existe usuario, carga sus datos desde Firestore. Si no existe, usa defaultDatabase como datos de demostracion."
    )
    add_code(
        doc,
        """
export function normalizeDatabase(savedDb) {
  const db = structuredClone(defaultDatabase);
  const incoming = savedDb && typeof savedDb === "object" ? savedDb : {};
  db.user = { ...db.user, ...(incoming.user || {}) };
  ...
  return db;
}
""",
    )
    doc.add_paragraph(
        "normalizeDatabase() evita errores cuando Firestore tiene datos incompletos o antiguos. Parte de la estructura por defecto y mezcla encima lo guardado. "
        "Esto permite agregar campos nuevos sin romper usuarios anteriores."
    )
    add_note(
        doc,
        "Por que se usa structuredClone",
        "Se usa para crear copias independientes de los datos base. Si se modificara defaultDatabase directamente, los datos de demostracion podrian contaminarse entre pantallas o sesiones.",
    )

    add_heading(doc, "7. firebase-service.js: capa de Firebase y Firestore")
    doc.add_paragraph(
        "Este archivo es la capa de comunicacion con Firebase. Si el profesor pregunta donde se conecta la app con la nube, la respuesta principal es firebase-service.js."
    )
    add_code(
        doc,
        """
import { initializeApp } from "https://www.gstatic.com/firebasejs/10.12.5/firebase-app.js";
import { getAuth, signInWithEmailAndPassword } from "https://www.gstatic.com/firebasejs/10.12.5/firebase-auth.js";
import { getFirestore, doc, getDoc, writeBatch } from "https://www.gstatic.com/firebasejs/10.12.5/firebase-firestore.js";
""",
    )
    doc.add_paragraph(
        "Estos imports vienen desde el CDN oficial de Firebase. No se instala Firebase con npm; el navegador descarga los modulos. "
        "firebase-app inicializa el proyecto, firebase-auth maneja usuarios y firebase-firestore maneja documentos."
    )
    add_code(
        doc,
        """
let app = null;
let auth = null;
let db = null;
""",
    )
    doc.add_paragraph(
        "Estas variables se declaran fuera de las funciones para que el modulo recuerde la instancia de Firebase. "
        "Asi initializeFirebase() no crea una app nueva cada vez que se llama."
    )
    add_code(
        doc,
        """
export async function loginWithEmail(email, password) {
  try {
    const credential = await signInWithEmailAndPassword(auth, email, password);
    return credential.user;
  } catch (error) {
    if (error.code !== "auth/user-not-found") throw error;
    const credential = await createUserWithEmailAndPassword(auth, email, password);
    return credential.user;
  }
}
""",
    )
    doc.add_paragraph(
        "La funcion autentica al usuario existente mediante Firebase Authentication. "
        "Esto facilita la demo, aunque en un sistema real podria separarse registro e inicio de sesion para tener mayor control."
    )
    add_code(
        doc,
        """
const [profileSnap, courseSnap, taskSnap] = await Promise.all([
  getDoc(userRef),
  getDocs(query(collection(db, "cursos"), where("userId", "==", user.uid))),
  getDocs(query(collection(db, "tareasAsistente"), where("userId", "==", user.uid)))
]);
""",
    )
    doc.add_paragraph(
        "Promise.all permite pedir perfil, cursos y tareas al mismo tiempo. Eso es mejor que esperar una consulta y luego la siguiente. "
        "Cada consulta de cursos y tareas filtra por userId, por eso cada estudiante ve su propia informacion."
    )
    add_code(
        doc,
        """
const batch = writeBatch(db);
batch.set(doc(db, "usuarios", userId), userData, { merge: true });
database.courses.forEach((course, index) => {
  batch.set(doc(db, "cursos", `${userId}_${course.id}`), courseDoc, { merge: true });
});
await batch.commit();
""",
    )
    doc.add_paragraph(
        "writeBatch agrupa varias escrituras y las envia juntas. En SARC se usa porque guardar el estado implica actualizar usuario, cursos, recomendaciones y tareas. "
        "El id del documento de curso se arma como userId_courseId para que sea facil ubicar el curso de un usuario especifico."
    )

    add_heading(doc, "8. navigation.js: rutas hash simples")
    add_code(
        doc,
        """
export function navigate(route) {
  window.location.hash = route;
}

export function openDetail(courseId) {
  state.currentDetailId = courseId;
  window.location.hash = `detalle/${courseId}`;
}
""",
    )
    doc.add_paragraph(
        "navigation.js resuelve un problema concreto: mover la aplicacion entre vistas sin recargar toda la pagina. "
        "Cuando cambia window.location.hash, app.js escucha el evento hashchange y vuelve a renderizar."
    )

    add_heading(doc, "9. courses.js: reglas de negocio de cursos")
    doc.add_paragraph(
        "courses.js no dibuja una vista completa. Se encarga de acciones compartidas sobre cursos: filtrar, buscar, inscribir y enlazar botones."
    )
    add_code(
        doc,
        """
export function getFilteredCourses() {
  return state.db.courses.filter((course) => {
    if (!CORE_COURSE_IDS.includes(course.id)) return false;
    const areaOk = !state.filters.area || state.filters.area === "Todas las areas" || course.area === state.filters.area;
    const modalityOk = !state.filters.modality || course.modality === state.filters.modality;
    return areaOk && modalityOk;
  });
}
""",
    )
    doc.add_paragraph(
        "La funcion filtra los cursos que se muestran en recomendaciones. Primero limita el catalogo a los cursos base. "
        "Luego verifica si el area coincide y si la modalidad coincide. Si no hay filtro, deja pasar el curso."
    )
    add_code(
        doc,
        """
course.enrolled = true;
course.registered = true;
course.progress = 5;
course.status = "En progreso";
course.lastAccess = "Hoy";
course.seats -= 1;
await saveDatabase();
refreshDashboard();
""",
    )
    doc.add_paragraph(
        "Este bloque de handleEnrollment() transforma un curso disponible en curso inscrito. Actualiza el objeto en memoria, guarda en Firestore, "
        "vuelve a pintar el dashboard y muestra retroalimentacion visual al usuario."
    )

    add_heading(doc, "10. template-loader.js y feedback.js: reutilizacion")
    doc.add_paragraph(
        "template-loader.js resuelve la carga de HTML externo. Cada vista tiene su archivo HTML y el JavaScript lo inserta en el contenedor. "
        "feedback.js resuelve mensajes reutilizables: modales, toasts y enlaces visuales."
    )
    add_code(
        doc,
        """
export async function loadTemplate(path) {
  if (templateCache.has(path)) return templateCache.get(path);
  const response = await fetch(path);
  const html = await response.text();
  templateCache.set(path, html);
  return html;
}
""",
    )
    doc.add_paragraph(
        "templateCache evita descargar la misma plantilla varias veces. La primera vez se usa fetch; despues se lee desde memoria."
    )
    add_code(
        doc,
        """
export function interpolate(template, values = {}) {
  return template.replace(/\\{\\{(\\w+)\\}\\}/g, (match, key) => {
    return Object.prototype.hasOwnProperty.call(values, key) ? String(values[key]) : "";
  });
}
""",
    )
    doc.add_paragraph(
        "interpolate reemplaza marcadores como {{name}} o {{courseList}} por datos reales. Es una solucion sencilla para plantillas, sin usar un framework."
    )

    add_heading(doc, "11. Como trabajan juntos HTML, CSS y JavaScript")
    doc.add_paragraph(
        "El patron se repite en casi todas las vistas. El HTML define la estructura y los espacios dinamicos. El CSS define la apariencia. "
        "El JavaScript carga la plantilla, reemplaza datos, inserta el resultado en el DOM y despues registra eventos."
    )
    number(doc, "HTML: contiene etiquetas, ids y clases. Ejemplo: un boton tiene id=\"editProfileButton\".")
    number(doc, "CSS: usa esas clases para dar color, espaciado, tarjetas, banners y distribucion.")
    number(doc, "JavaScript: busca el boton con querySelector y le agrega addEventListener para responder al clic.")
    add_note(
        doc,
        "Ejemplo oral",
        "En perfil.html el boton Editar Perfil solo existe como estructura. En perfil.css se ve como boton. En perfil.js se vuelve funcional porque bindProfileEvents() le agrega el evento click que abre el modal.",
        LIGHT_GREEN,
    )

    add_heading(doc, "12. Vista Login: autenticacion")
    doc.add_paragraph(
        "login.js carga login.html, enlaza el formulario y captura el submit. La funcion handleLogin() evita que el formulario recargue la pagina, "
        "lee correo y contrasena, llama a Firebase, carga los datos del usuario y navega al inicio."
    )
    add_code(
        doc,
        """
const email = document.getElementById("loginEmail").value.trim().toLowerCase();
const password = document.getElementById("loginPassword").value.trim();
const user = await loginWithEmail(email, password);
await loadDatabaseForUser(user);
saveSession(loadSession(user));
navigate("inicio");
""",
    )
    doc.add_paragraph(
        "trim() limpia espacios accidentales. toLowerCase() normaliza el correo. await se usa porque Firebase responde de forma asincrona. "
        "Si hay error, se muestra un modal con un mensaje entendible."
    )

    add_heading(doc, "13. Vista Inicio: cursos sugeridos y asistente")
    doc.add_paragraph(
        "inicio.js muestra una bienvenida, botones de cursos sugeridos y tareas del asistente. Los botones de curso no se escriben fijos en el HTML: "
        "se generan recorriendo state.db.courses."
    )
    add_code(
        doc,
        """
function renderCourseButtons(courses) {
  return courses
    .map((course) => `<button type="button" data-detail="${course.id}">○ ${course.name}</button>`)
    .join("");
}
""",
    )
    doc.add_paragraph(
        "map convierte cada curso en un boton HTML. join une todos los botones en un solo texto. data-detail guarda el id del curso para que bindCourseActions() sepa que detalle abrir."
    )

    add_heading(doc, "14. Vista Recomendaciones: catalogo con filtros")
    doc.add_paragraph(
        "recomendaciones.js combina el estado de filtros con la lista de cursos. Cuando cambia un select, actualiza state.filters y vuelve a llamar renderRecommendations(container). "
        "Asi la pantalla se redibuja con los nuevos resultados."
    )
    add_code(
        doc,
        """
container.querySelector("#filterArea").addEventListener("change", (event) => {
  state.filters.area = event.target.value;
  renderRecommendations(container);
});
""",
    )
    doc.add_paragraph(
        "El evento change se dispara cuando el usuario selecciona otra opcion. event.target.value es el valor seleccionado. "
        "La vista se renderiza otra vez para reflejar el filtro."
    )

    add_heading(doc, "15. Vista Detalle: curso seleccionado")
    doc.add_paragraph(
        "detalle.js depende de state.currentDetailId. Ese valor se establece cuando el usuario pulsa Ver detalles. "
        "Luego getCourseById() busca el objeto completo del curso y la plantilla muestra descripcion, duracion, nivel y boton de inscripcion."
    )
    add_code(
        doc,
        """
const course = getCourseById(state.currentDetailId) || state.db.courses[0];
container.innerHTML = interpolate(template, {
  courseId: course.id,
  courseName: course.name,
  description: course.description
});
""",
    )
    doc.add_paragraph(
        "El operador || usa un curso por defecto si no se encuentra el id. Esto evita que la pantalla quede rota por una ruta invalida."
    )

    add_heading(doc, "16. Vista Progreso: barras y reportes")
    doc.add_paragraph(
        "progreso.js filtra los cursos inscritos, crea filas de progreso y permite abrir un reporte detallado. "
        "La animacion de barras se hace despues de pintar el DOM usando requestAnimationFrame()."
    )
    add_code(
        doc,
        """
requestAnimationFrame(() => {
  container.querySelectorAll(".progress-fill[data-width]").forEach((bar) => {
    bar.style.width = bar.dataset.width;
  });
});
""",
    )
    doc.add_paragraph(
        "data-width guarda el porcentaje como dato HTML. Luego JavaScript lo copia a style.width. "
        "Esto permite que CSS anime el cambio de ancho de la barra."
    )
    doc.add_paragraph(
        "Cuando el usuario pulsa Descargar, progreso.js llama downloadReport(course). Esa funcion esta en pdf.js y construye un PDF simple en bytes, sin depender de librerias externas."
    )

    add_heading(doc, "17. Vista Perfil: edicion y cierre de sesion")
    doc.add_paragraph(
        "perfil.js muestra los datos del estudiante desde state.db.user. Tambien abre un modal para cambiar modalidad, avatar y contrasena. "
        "El cambio de contrasena se hace con Firebase Authentication, no guardando contrasenas en Firestore."
    )
    add_code(
        doc,
        """
if (newPassword) {
  await updateCurrentUserPassword(currentPassword, newPassword);
}
state.db.user.modality = selectedModality;
await saveDatabase();
""",
    )
    doc.add_paragraph(
        "La modalidad se guarda en Firestore porque es dato del perfil. La contrasena se actualiza con Firebase Auth porque es informacion sensible. "
        "Esta separacion es importante para defender seguridad basica del proyecto."
    )

    add_heading(doc, "18. Flujo entre frontend, Firebase y Firestore")
    number(doc, "La vista recibe una accion del usuario: login, inscripcion, cambio de modalidad o tarea marcada.")
    number(doc, "La vista actualiza state.db o pide autenticacion mediante una funcion importada.")
    number(doc, "data.js decide si debe guardar y llama saveUserData() de firebase-service.js.")
    number(doc, "firebase-service.js escribe documentos en usuarios, cursos, recomendaciones y tareasAsistente.")
    number(doc, "Firestore valida las reglas. Si request.auth.uid coincide con userId, permite la operacion.")
    number(doc, "La vista se vuelve a pintar para mostrar el resultado actualizado.")
    add_table(
        doc,
        ["Coleccion", "Que guarda"],
        [
            ["usuarios", "Perfil del estudiante: nombre, correo, facultad, carrera, semestre, modalidad y avatar."],
            ["cursos", "Estado de cada curso por usuario: progreso, cupos, inscripcion, estado y metricas."],
            ["recomendaciones", "Recomendacion y alternativas asociadas a cada curso del usuario."],
            ["tareasAsistente", "Checklist del asistente virtual, tambien separado por userId."],
        ],
        widths=[1.5, 4.8],
    )

    add_heading(doc, "19. Reglas de Firestore")
    doc.add_paragraph(
        "firestore.rules es la defensa del lado de la base de datos. Aunque alguien modifique JavaScript en el navegador, Firestore solo permite operaciones si el usuario autenticado coincide con el documento."
    )
    add_code(
        doc,
        """
function ownsResource() {
  return signedIn() && request.auth.uid == resource.data.userId;
}

function ownsIncomingResource() {
  return signedIn() && request.auth.uid == request.resource.data.userId;
}
""",
    )
    doc.add_paragraph(
        "ownsResource() se usa para leer, actualizar o borrar documentos existentes. ownsIncomingResource() se usa para crear documentos nuevos. "
        "La diferencia es que resource.data mira lo que ya existe y request.resource.data mira lo que se quiere guardar."
    )

    add_heading(doc, "20. Buenas practicas que puedes mencionar")
    bullet(doc, "Uso de modulos ES: cada archivo exporta funciones concretas e importa lo que necesita.")
    bullet(doc, "Separacion por vistas: HTML, CSS y JS estan organizados por pantalla.")
    bullet(doc, "Capa de servicio para Firebase: la UI no repite codigo de Firestore en cada vista.")
    bullet(doc, "Uso de userId para separar datos de usuarios.")
    bullet(doc, "Uso de writeBatch para guardar varias colecciones de forma ordenada.")
    bullet(doc, "Validaciones de formulario antes de enviar cambios.")
    bullet(doc, "Reglas de Firestore que complementan la proteccion del frontend.")

    add_heading(doc, "21. Errores o mejoras posibles")
    doc.add_paragraph(
        "Estas observaciones no destruyen el proyecto; al contrario, muestran que entiendes sus limites y sabes como evolucionarlo."
    )
    bullet(doc, "Separar registro e inicio de sesion. Actualmente loginWithEmail crea el usuario si no existe, util para demo pero menos controlado en produccion.")
    bullet(doc, "No reiniciar siempre la sesion al abrir. resetStartupSession() sirve para la demostracion, pero una app real podria conservar la sesion del usuario.")
    bullet(doc, "Validar y sanitizar mejor HTML dinamico. interpolate y los templates son simples; si entraran textos libres de usuarios, convendria evitar insertar HTML sin control.")
    bullet(doc, "Agregar manejo visual de estados de carga y errores en cada vista, no solo en login.")
    bullet(doc, "Subir avatar a Firebase Storage. Ahora el avatar se guarda como data URL en Firestore, suficiente para demo pero no ideal para imagenes grandes.")
    bullet(doc, "Agregar pruebas o checklist funcional para login, inscripcion, filtros, progreso y perfil.")

    add_heading(doc, "22. Preguntas tipicas del profesor y respuestas")
    qa = [
        (
            "Por que no usaste React o Angular?",
            "Porque el objetivo era construir una app web entendible con tecnologias base. JavaScript modular permite separar archivos, usar imports/exports y mantener el proyecto ordenado sin framework.",
        ),
        (
            "Donde se guardan los datos?",
            "Los datos se guardan en Cloud Firestore. En memoria la app usa state.db, pero cuando hay cambios importantes se llama saveDatabase() y eso termina en saveUserData() dentro de firebase-service.js.",
        ),
        (
            "Como evitas que un usuario vea datos de otro?",
            "Cada documento tiene userId y las reglas de Firestore comparan request.auth.uid contra ese userId. Ademas, las consultas filtran por userId.",
        ),
        (
            "Que hace app.js?",
            "Es el coordinador de la aplicacion: inicializa estado, registra eventos globales, lee la ruta hash, protege vistas privadas y llama la funcion render de cada pantalla.",
        ),
        (
            "Que hace data.js?",
            "Mantiene el estado global, define los datos base, carga datos del usuario, normaliza informacion guardada y ofrece saveDatabase() para persistir cambios.",
        ),
        (
            "Que hace firebase-service.js?",
            "Encapsula Firebase Authentication y Firestore: login, recuperacion, cambio de contrasena, carga de datos, guardado por batch y reseteo de datos.",
        ),
        (
            "Como funciona una inscripcion?",
            "El boton tiene data-enroll con el id del curso. bindCourseActions() conecta el clic con handleEnrollment(). Esa funcion busca el curso, valida cupos, cambia sus propiedades, guarda en Firestore y refresca la pantalla.",
        ),
        (
            "Como se generan las vistas?",
            "Cada render carga una plantilla HTML con loadTemplate(), reemplaza marcadores con interpolate(), inserta el HTML en el contenedor y luego enlaza eventos.",
        ),
    ]
    for question, answer in qa:
        p = doc.add_paragraph()
        r = p.add_run(f"Pregunta: {question}")
        r.bold = True
        r.font.color.rgb = BLUE
        doc.add_paragraph(f"Respuesta: {answer}")

    add_heading(doc, "23. Guion recomendado para la demostracion")
    number(doc, "Abrir la app en localhost y explicar que app.py solo sirve archivos estaticos.")
    number(doc, "Mostrar index.html y ubicar authView, dashboardView y viewContainer.")
    number(doc, "Explicar app.js: init(), renderApp() y pageRenderers.")
    number(doc, "Iniciar sesion y explicar login.js + loginWithEmail().")
    number(doc, "Ir a recomendaciones, aplicar filtros y mostrar getFilteredCourses().")
    number(doc, "Inscribirse a un curso y explicar handleEnrollment() + saveDatabase().")
    number(doc, "Abrir Firestore y mostrar documentos con userId.")
    number(doc, "Ir a progreso, abrir reporte y explicar pdf.js de forma general.")
    number(doc, "Ir a perfil, cambiar modalidad y explicar que contrasena va por Firebase Auth.")
    number(doc, "Cerrar con reglas de Firestore y mejoras posibles.")

    add_heading(doc, "24. Cierre para defender con seguridad")
    doc.add_paragraph(
        "La idea central que debes repetir es que SARC no es solo una maqueta visual: tiene una estructura modular, un flujo de datos claro y persistencia real. "
        "HTML define las pantallas, CSS les da presentacion, JavaScript controla comportamiento y Firebase guarda/autentica. "
        "Si mantienes esa linea de explicacion, puedes responder la mayoria de preguntas tecnicas con seguridad."
    )

    polish_spanish_accents(doc)
    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
