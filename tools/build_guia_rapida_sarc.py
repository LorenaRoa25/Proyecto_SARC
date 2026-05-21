"""Genera la Guia Rapida de Uso del sistema SARC."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parents[1]
OUT = PROJECT_DIR / "Documentos" / "Guia_Rapida_Uso_SARC.docx"

BLUE = "0B4F8A"
GREEN = "6AA84F"
LIGHT_BLUE = "EAF3FA"
LIGHT_GREEN = "EAF5E6"
LIGHT_YELLOW = "FFF4CC"
LIGHT_RED = "FDECEC"
TEXT = RGBColor(31, 41, 55)


def set_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def configure(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(3)
        section.left_margin = Cm(4)
        section.right_margin = Cm(2)
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Pagina ")
        r.font.name = "Arial"
        r.font.size = Pt(9)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        p._p.append(fld)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(2)

    for name in ["Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(2)
    doc.styles["Heading 1"].font.size = Pt(12)
    doc.styles["Heading 2"].font.size = Pt(10)


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = TEXT


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(item)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = TEXT


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(item)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = TEXT


def callout(doc: Document, label: str, text: str, kind: str = "note") -> None:
    fill = {"note": LIGHT_BLUE, "important": LIGHT_YELLOW, "warning": LIGHT_RED}.get(kind, LIGHT_BLUE)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_shading(cell, fill)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.color.rgb = TEXT


def small_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        set_shading(cell, BLUE)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            p = cell.paragraphs[0]
            r = p.add_run(value)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = TEXT
    doc.add_paragraph()


def cover(doc: Document) -> None:
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GUIA RAPIDA DE USO DEL SISTEMA DE APOYO ACADEMICO CON RECOMENDACION DE CURSOS (SARC)")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    doc.add_paragraph()
    for line in [
        "Guia para estudiantes usuarios del sistema",
        "Corporacion Universitaria de Sabaneta - Unisabaneta",
        "Proyecto SARC",
        "Sabaneta, Antioquia",
        "2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.name = "Arial"
        r.font.size = Pt(12)
    doc.add_page_break()


def build_doc() -> Document:
    doc = Document()
    configure(doc)
    cover(doc)

    heading(doc, "1. Introduccion breve")
    para(doc, "SARC es una plataforma web academica que ayuda a los estudiantes a consultar cursos complementarios, recibir recomendaciones, inscribirse en cursos disponibles y revisar su progreso academico. Su objetivo principal es facilitar el acceso a oportunidades de formacion y apoyar el fortalecimiento de competencias.")
    callout(doc, "Nota", "Las recomendaciones del sistema se basan en reglas, cursos inscritos, progreso y rendimiento. No es un chat de inteligencia artificial.", "note")

    heading(doc, "2. Requisitos para usar el sistema")
    bullets(doc, [
        "Navegador recomendado: Google Chrome o Microsoft Edge actualizado.",
        "Conexion a internet estable para usar Firebase Authentication y Firestore.",
        "Correo institucional y contrasena registrados en el sistema.",
        "Computador, portatil, tablet o celular con navegador moderno.",
    ])

    heading(doc, "3. Como iniciar sesion")
    numbered(doc, [
        "Si usa la aplicacion instalada, haga doble clic en Proyecto SARC; si usa una version web, abra la direccion publicada.",
        "Escriba su correo institucional en el campo Correo institucional.",
        "Escriba su contrasena.",
        "Acepte el tratamiento de datos personales.",
        "Presione el boton INGRESAR para acceder al panel principal.",
    ])
    callout(doc, "Importante", "Use contrasenas seguras y no comparta sus credenciales con otras personas.", "important")

    heading(doc, "4. Como recuperar contrasena")
    numbered(doc, [
        "Seleccione la opcion Recuperar contrasena en la pantalla de inicio de sesion.",
        "Ingrese su correo institucional.",
        "Escriba y confirme la nueva contrasena solicitada por la pantalla.",
        "Revise las instrucciones enviadas por Firebase al correo registrado.",
        "Regrese al inicio de sesion e ingrese nuevamente.",
    ])
    callout(doc, "Advertencia", "Si no recibe el correo, verifique que la cuenta exista en Firebase y revise la carpeta de spam.", "warning")

    heading(doc, "5. Panel principal")
    para(doc, "Al ingresar, el sistema muestra el dashboard con el saludo del estudiante, cursos sugeridos y un panel de Asistente Virtual. Desde la barra superior puede navegar a Inicio, Recomendaciones, Mi progreso y Perfil.")
    bullets(doc, [
        "Inicio: muestra cursos sugeridos y recomendaciones generales.",
        "Recomendaciones: permite explorar cursos y aplicar filtros.",
        "Mi progreso: muestra avance, actividades, nota final y reportes.",
        "Perfil: muestra datos personales, modalidad, avatar y acciones de cuenta.",
    ])

    heading(doc, "6. Consultar cursos recomendados")
    numbered(doc, [
        "Ingrese al modulo Recomendaciones desde la barra superior.",
        "Revise la lista de cursos sugeridos.",
        "Use los filtros de area interdisciplinaria y modalidad si desea reducir la lista.",
        "Presione Ver detalles para consultar descripcion, duracion y nivel del curso.",
    ])

    heading(doc, "7. Inscribirse a un curso")
    numbered(doc, [
        "Seleccione un curso desde Recomendaciones o desde el detalle del curso.",
        "Revise la informacion disponible del curso.",
        "Presione INSCRIBIRSE.",
        "Si hay cupos, el sistema mostrara confirmacion de inscripcion.",
        "Puede volver al Inicio o ir a Mi progreso.",
    ])
    callout(doc, "Nota", "Si el curso no tiene cupos, el sistema muestra un mensaje con cursos alternativos configurados.", "note")

    heading(doc, "8. Consultar progreso academico")
    para(doc, "En Mi progreso se visualizan los cursos inscritos, porcentaje de avance, estado, ultimo acceso, actividades evaluativas, calculo de nota final y rendimiento academico. Tambien se puede abrir un reporte detallado y descargar un PDF simple.")

    heading(doc, "9. Consultar perfil")
    para(doc, "En Perfil se muestran nombre, facultad, carrera, semestre, modalidad preferida y foto. Desde Editar Perfil puede cambiar modalidad, actualizar contrasena y seleccionar una imagen de perfil. Tambien puede cerrar sesion o reiniciar la demo.")

    heading(doc, "10. Asistente virtual")
    para(doc, "El Asistente Virtual aparece en el panel principal y muestra recomendaciones academicas basadas en el progreso y rendimiento del estudiante. En la version actual no permite escribir preguntas; funciona como un panel de orientacion automatica.")

    heading(doc, "11. Recomendaciones y buenas practicas")
    bullets(doc, [
        "Revise periodicamente los cursos sugeridos.",
        "Mantenga actualizada su modalidad preferida en el perfil.",
        "Use contrasenas fuertes y cierre sesion al terminar.",
        "Consulte Mi progreso para conocer su estado academico.",
        "Descargue reportes cuando necesite guardar evidencia de avance.",
    ])

    heading(doc, "12. Solucion de problemas frecuentes")
    small_table(doc, ["Problema", "Solucion rapida"], [
        ["No puedo ingresar", "Verifique correo, contrasena y aceptacion de datos personales."],
        ["No carga el sistema", "Abra nuevamente Proyecto SARC, recargue con Ctrl + F5 y confirme conexion a internet."],
        ["No aparecen cursos", "Revise filtros de area/modalidad o vuelva a Inicio."],
        ["No llega correo de recuperacion", "Revise spam y confirme que el correo este registrado."],
        ["La navegacion no cambia", "Use la barra superior o recargue la pagina del sistema."],
    ], [2.0, 4.5])

    heading(doc, "13. Preguntas frecuentes y glosario")
    small_table(doc, ["Pregunta / Termino", "Respuesta breve"], [
        ["Que es una recomendacion", "Sugerencia de curso o accion academica generada por el sistema."],
        ["Que es dashboard", "Panel principal donde se muestran cursos sugeridos y asistente."],
        ["Que es perfil", "Seccion con datos personales, modalidad y opciones de cuenta."],
        ["Que es inscripcion", "Accion de registrar un curso como parte del progreso del estudiante."],
        ["Puedo usar IA conversacional", "No. El asistente actual muestra recomendaciones, no es un chat."],
    ], [2.0, 4.5])

    heading(doc, "14. Conclusion")
    para(doc, "SARC facilita el acceso a cursos complementarios y ayuda al estudiante a revisar recomendaciones, inscribirse y hacer seguimiento de su avance academico. Su uso frecuente puede apoyar mejores decisiones de formacion y fortalecer competencias importantes para la vida universitaria.")

    core = doc.core_properties
    core.title = "Guia Rapida de Uso SARC"
    core.subject = "Guia rapida para estudiantes usuarios del Sistema SARC"
    core.author = "Lorena Roa Rivera"
    core.keywords = "SARC, guia de uso, estudiantes, cursos, recomendaciones"
    return doc


def main() -> None:
    doc = build_doc()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
