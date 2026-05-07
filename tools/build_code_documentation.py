"""Genera la documentacion tecnica del codigo SARC en formato DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_DIR / "DOCUMENTACION_CODIGO_SARC.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document, text, level=1):
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(8, 83, 148)
    return heading


def add_code_block(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    return paragraph


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for index, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[index], "085394")
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=(255, 255, 255))

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)

    document.add_paragraph()
    return table


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Segoe UI"
    styles["Normal"].font.size = Pt(10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Documentacion del Codigo\nSARC")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(8, 83, 148)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Sistema de Apoyo Academico con Recomendacion de Cursos")
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = RGBColor(80, 80, 80)

    add_heading(document, "1. Resumen del proyecto", 1)
    document.add_paragraph(
        "SARC es una aplicacion web local construida con HTML5, CSS3 y JavaScript puro. "
        "Implementa autenticacion con Firebase, recuperacion de contrasena, catalogo de cursos, "
        "inscripcion, progreso del estudiante, reportes PDF y edicion de perfil."
    )

    add_heading(document, "2. Tecnologias utilizadas", 1)
    add_table(
        document,
        ["Tecnologia", "Uso en el proyecto"],
        [
            ["HTML5", "Estructura de vistas y plantillas de pantalla."],
            ["CSS3", "Estilos globales, componentes y estilos por vista."],
            ["JavaScript puro", "Rutas hash, estado, eventos, modales, Firebase y reportes."],
            ["Firebase Authentication", "Login y recuperacion con correo y contrasena."],
            ["Cloud Firestore", "Persistencia por userId en colecciones globales."],
            ["Python", "Servidor local simple para ejecutar el proyecto desde el navegador."],
        ],
    )

    add_heading(document, "3. Estructura de carpetas", 1)
    add_code_block(
        document,
        """Proyecto_SARC_1/
  index.html
  app.py
  README.txt
  FIREBASE_SETUP.md
  firestore.rules
  DOCUMENTACION_CODIGO_SARC.docx
  shared/
    css/
    js/
    components/
  pages/
    login/
    recuperar/
    inicio/
    recomendaciones/
    detalle/
    progreso/
    perfil/
  tools/""",
    )

    add_heading(document, "4. Archivos principales", 1)
    add_table(
        document,
        ["Archivo", "Responsabilidad"],
        [
            ["index.html", "Shell principal, enlaces CSS y entrada JavaScript."],
            ["shared/js/app.js", "Inicializa estado, controla rutas y renderiza vistas."],
            ["shared/js/data.js", "Define datos base, estado global y sincronizacion por usuario."],
            ["shared/js/firebase-service.js", "Centraliza Auth, lecturas, escrituras y reseteo en Firestore."],
            ["shared/js/firebase-config.js", "Configuracion del proyecto Firebase."],
            ["shared/js/courses.js", "Filtros, busqueda e inscripcion de cursos."],
            ["shared/components/feedback.js", "Modales, toasts y enlaces visuales del prototipo."],
            ["firestore.rules", "Reglas de seguridad para aislar datos por userId."],
            ["app.py", "Servidor HTTP local simple para desarrollo."],
        ],
    )

    add_heading(document, "5. Vistas del sistema", 1)
    add_table(
        document,
        ["Vista", "Archivos", "Descripcion"],
        [
            ["Login", "pages/login/login.*", "Autentica con Firebase y carga datos del usuario."],
            ["Recuperar", "pages/recuperar/recuperar.*", "Envia recuperacion de contrasena por Firebase."],
            ["Inicio", "pages/inicio/inicio.*", "Muestra cursos sugeridos y asistente virtual."],
            ["Recomendaciones", "pages/recomendaciones/recomendaciones.*", "Catalogo filtrable e inscripcion."],
            ["Detalle", "pages/detalle/detalle.*", "Descripcion, duracion, nivel e inscripcion."],
            ["Progreso", "pages/progreso/progreso.*", "Avances y reportes detallados."],
            ["Perfil", "pages/perfil/perfil.*", "Avatar, modalidad, contrasena y sesion."],
        ],
    )

    add_heading(document, "6. Flujo de ejecucion", 1)
    for step in [
        "El navegador carga index.html.",
        "index.html importa shared/js/app.js como modulo ES.",
        "app.js inicializa Firebase y el estado global.",
        "La ruta hash define si se renderiza autenticacion o dashboard.",
        "Cada vista carga su HTML con template-loader.js y enlaza sus eventos.",
        "Los cambios se guardan en Firestore filtrados por userId.",
    ]:
        document.add_paragraph(step, style="List Number")

    add_heading(document, "7. Funciones importantes", 1)
    add_table(
        document,
        ["Funcion", "Ubicacion", "Proposito"],
        [
            ["renderApp()", "shared/js/app.js", "Decide que vista mostrar segun ruta y sesion."],
            ["initializeState()", "shared/js/data.js", "Inicializa Firebase, sesion y datos del usuario."],
            ["loginWithEmail(email, password)", "shared/js/firebase-service.js", "Autentica o registra usuarios."],
            ["saveUserData(userId, database)", "shared/js/firebase-service.js", "Persiste perfil, cursos, recomendaciones y tareas."],
            ["handleEnrollment(courseId)", "shared/js/courses.js", "Gestiona inscripcion, cupos, mensajes y persistencia."],
            ["showModal(config)", "shared/components/feedback.js", "Renderiza modales reutilizables."],
            ["downloadReport(course)", "shared/js/pdf.js", "Genera y descarga PDF de progreso."],
        ],
    )

    add_heading(document, "8. Persistencia Firebase", 1)
    document.add_paragraph(
        "El proyecto usa Firebase Authentication para usuarios y Cloud Firestore para datos. "
        "Las colecciones globales son usuarios, cursos, recomendaciones y tareasAsistente. "
        "Cada documento de datos incluye userId, lo que permite separar multiples usuarios sin "
        "crear una coleccion distinta para cada uno."
    )

    add_heading(document, "9. Detalle tecnico que conviene sustentar", 1)
    document.add_paragraph(
        "La documentacion tecnica no debe quedarse solo en el listado de archivos. Para defender el proyecto, "
        "tambien es importante explicar como circulan los datos entre las vistas, el estado global y Firebase."
    )

    add_heading(document, "9.1 Flujo interno del estado", 2)
    for step in [
        "Las vistas no consultan directamente Firestore; leen y modifican state.db, definido en shared/js/data.js.",
        "Cuando una accion cambia datos importantes, como una inscripcion o una tarea del asistente, se llama saveDatabase().",
        "saveDatabase() obtiene el usuario autenticado y delega la escritura a saveUserData() en firebase-service.js.",
        "firebase-service.js guarda usuario, cursos, recomendaciones y tareas en colecciones separadas, todas asociadas a userId.",
    ]:
        document.add_paragraph(step, style="List Bullet")

    add_heading(document, "9.2 Enrutamiento y vistas", 2)
    document.add_paragraph(
        "La navegacion se basa en el hash de la URL. navigation.js cambia window.location.hash y app.js escucha el evento "
        "hashchange para decidir que vista renderizar. Si la ruta empieza por detalle/, app.js separa el identificador del curso "
        "y lo guarda en state.currentDetailId."
    )

    add_heading(document, "9.3 Plantillas HTML", 2)
    document.add_paragraph(
        "Cada vista carga su plantilla HTML mediante loadTemplate(). Luego interpolate() reemplaza marcadores como {{courseList}} "
        "o {{name}} por datos reales. Despues de insertar el HTML en el DOM, la vista registra eventos con addEventListener()."
    )

    add_heading(document, "9.4 Firebase y Firestore", 2)
    document.add_paragraph(
        "Firebase Authentication administra credenciales y sesiones. Firestore guarda los datos academicos del usuario. "
        "El proyecto no guarda contrasenas en Firestore. Las reglas de seguridad comparan request.auth.uid con userId para "
        "evitar acceso cruzado entre usuarios."
    )

    add_heading(document, "9.5 Aspectos a mejorar", 2)
    for item in [
        "Separar formalmente registro e inicio de sesion; loginWithEmail crea el usuario si no existe, lo cual es util para demo.",
        "No reiniciar siempre la sesion al cargar si se desea comportamiento productivo persistente.",
        "Subir avatares a Firebase Storage en vez de guardar imagenes como data URL dentro del perfil.",
        "Agregar pruebas funcionales para login, filtros, inscripcion, progreso, perfil y reglas de Firestore.",
    ]:
        document.add_paragraph(item, style="List Bullet")

    add_heading(document, "10. Limpieza aplicada", 1)
    for item in [
        "Eliminacion de carpetas vacias de assets.",
        "Eliminacion del modulo dashboard reservado que no estaba referenciado.",
        "Retiro de constantes antiguas de localStorage y codigo de depuracion.",
        "Actualizacion del README y documentacion tecnica a Firebase.",
        "Login sin credenciales precargadas.",
    ]:
        document.add_paragraph(item, style="List Bullet")

    add_heading(document, "11. Ejecucion recomendada", 1)
    add_code_block(document, "python app.py\nhttp://localhost:8000/index.html")
    document.add_paragraph(
        "Los usuarios se gestionan desde Firebase Authentication. El formulario de login no incluye credenciales precargadas."
    )

    document.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
