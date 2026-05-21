"""Genera el Manual de Usuario de SARC con capturas reales anotadas."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


PROJECT_DIR = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_DIR / "Documentos"
SHOT_DIR = DOCS_DIR / "manual_usuario_sarc_capturas"
OUT = DOCS_DIR / "Manual_Usuario_SARC.docx"

BLUE = "0B4F8A"
GREEN = "6AA84F"
LIGHT_BLUE = "EAF3FA"
LIGHT_GREEN = "EAF5E6"
LIGHT_GRAY = "F3F6F9"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(91, 103, 120)
WHITE = RGBColor(255, 255, 255)


def font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str) -> None:
    draw.line([start, end], fill=color, width=4)
    x1, y1 = start
    x2, y2 = end
    dx = x2 - x1
    dy = y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    left = (int(x2 - 16 * ux - 8 * uy), int(y2 - 16 * uy + 8 * ux))
    right = (int(x2 - 16 * ux + 8 * uy), int(y2 - 16 * uy - 8 * ux))
    draw.polygon([end, left, right], fill=color)


def label(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, fill: str = "#0B4F8A") -> None:
    fnt = font(17, bold=True)
    x, y = xy
    bbox = draw.textbbox((x, y), text, font=fnt)
    pad = 8
    draw.rounded_rectangle(
        (bbox[0] - pad, bbox[1] - pad, bbox[2] + pad, bbox[3] + pad),
        radius=8,
        fill=fill,
    )
    draw.text((x, y), text, fill="white", font=fnt)


def annotate_image(name: str, boxes: list[tuple[int, int, int, int, str]], arrows: list[tuple[tuple[int, int], tuple[int, int]]]) -> Path:
    raw = SHOT_DIR / f"{name}_raw.png"
    out = SHOT_DIR / f"{name}_anotada.png"
    img = Image.open(raw).convert("RGB")
    draw = ImageDraw.Draw(img)
    accent = "#0B4F8A"
    green = "#6AA84F"
    for box in boxes:
        x1, y1, x2, y2, title = box[:5]
        label_xy = box[5] if len(box) > 5 else (x1, max(8, y1 - 34))
        color = green if title.startswith("Boton") or title.startswith("Reporte") else accent
        draw.rounded_rectangle((x1, y1, x2, y2), radius=10, outline=color, width=5)
        label(draw, label_xy, title, color)
    for start, end in arrows:
        arrow(draw, start, end, "#0B4F8A")
    img.save(out, quality=95)
    return out


def build_annotations() -> dict[str, Path]:
    return {
        "login": annotate_image(
            "01_login",
            [
                (388, 236, 860, 291, "Correo", (276, 250)),
                (388, 305, 860, 358, "Contraseña", (246, 318)),
                (812, 312, 850, 350, "Ver contraseña", (862, 319)),
                (391, 374, 860, 414, "Habeas Data", (238, 388)),
                (388, 430, 860, 488, "Ingresar", (288, 449)),
            ],
            [((625, 505), (625, 488))],
        ),
        "inicio": annotate_image(
            "02_inicio",
            [
                (45, 210, 595, 585, "Cursos sugeridos"),
                (675, 210, 1235, 610, "Asistente virtual"),
            ],
            [((625, 360), (690, 360))],
        ),
        "recomendaciones": annotate_image(
            "03_recomendaciones",
            [
                (90, 184, 372, 230, "Filtros"),
                (38, 267, 1004, 414, "Info del curso"),
                (1010, 286, 1189, 390, "Botones"),
            ],
            [((372, 208), (598, 208)), ((920, 345), (1010, 345))],
        ),
        "progreso": annotate_image(
            "04_progreso",
            [
                (170, 215, 950, 305, "Avance general"),
                (171, 322, 948, 392, "Indicadores"),
                (970, 316, 1195, 366, "Reporte PDF"),
            ],
            [((890, 345), (970, 345))],
        ),
        "perfil": annotate_image(
            "05_perfil",
            [
                (90, 205, 575, 435, "Datos del estudiante"),
                (923, 205, 1139, 411, "Avatar"),
                (337, 527, 911, 595, "Acciones"),
            ],
            [((575, 320), (923, 310))],
        ),
    }


def set_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def configure(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Manual de Usuario SARC | Pagina ")
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        p._p.append(field)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.space_after = Pt(3)

    for style_name in ["Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(3)
    doc.styles["Heading 1"].font.size = Pt(12)
    doc.styles["Heading 2"].font.size = Pt(10)


def add_page_number_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def paragraph(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(BLUE)
        text = text[len(bold_prefix):]
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.color.rgb = TEXT


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(item)
        r.font.name = "Arial"
        r.font.size = Pt(8.7)
        r.font.color.rgb = TEXT


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(item)
        r.font.name = "Arial"
        r.font.size = Pt(8.7)
        r.font.color.rgb = TEXT


def callout(doc: Document, label_text: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_shading(cell, fill)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(f"{label_text}: ")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    r = p.add_run(body)
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.font.color.rgb = TEXT


def metadata_table(doc: Document) -> None:
    rows = [
        ("Sistema", "SARC - Sistema de Apoyo Academico con Recomendacion de Cursos"),
        ("Documento", "Manual de Usuario"),
        ("Dirigido a", "Estudiantes usuarios de la plataforma"),
        ("Universidad", "Corporacion Universitaria de Sabaneta - Unisabaneta"),
        ("Integrante", "Lorena Roa Rivera"),
        ("Fecha", "2026"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.55)
        cells[1].width = Inches(4.7)
        set_shading(cells[0], BLUE)
        set_cell_margins(cells[0])
        set_cell_margins(cells[1])
        r = cells[0].paragraphs[0].add_run(left)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = WHITE
        r = cells[1].paragraphs[0].add_run(right)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = TEXT


def figure(doc: Document, image_path: Path, caption: str, note: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(10.6))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    paragraph(doc, note, "Captura: ")


def cover(doc: Document) -> None:
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MANUAL DE USUARIO DEL SISTEMA DE APOYO ACADÉMICO CON RECOMENDACIÓN DE CURSOS (SARC)")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Guía práctica para estudiantes usuarios del sistema")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.color.rgb = TEXT

    doc.add_paragraph()
    cover_lines = [
        "Lorena Roa Rivera",
        "Ingeniería Informática",
        "Corporación Universitaria de Sabaneta - Unisabaneta",
        "Sabaneta, Antioquia",
        "2026",
    ]
    for line in cover_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.color.rgb = TEXT

    doc.add_paragraph()
    add_page_number_break(doc)


def build_doc(images: dict[str, Path]) -> None:
    doc = Document()
    configure(doc)
    cover(doc)

    doc.add_heading("1. Introduccion", level=1)
    paragraph(
        doc,
        "SARC es una plataforma academica que permite al estudiante consultar cursos recomendados, inscribirse en cursos disponibles, revisar su progreso academico y recibir orientacion basica mediante un asistente virtual. El sistema busca apoyar el seguimiento individual y facilitar la toma de decisiones academicas.",
    )

    doc.add_heading("2. Requisitos basicos", level=1)
    bullets(
        doc,
        [
            "Navegador actualizado, preferiblemente Google Chrome o Microsoft Edge.",
            "Conexion a internet para validar el acceso y cargar informacion desde Firebase.",
            "Correo institucional registrado en el sistema SARC.",
            "Aceptar el tratamiento de datos personales antes de ingresar.",
        ],
    )

    doc.add_heading("3. Acceso al sistema", level=1)
    numbered(
        doc,
        [
            "Abra la direccion local o web donde se encuentre publicado SARC.",
            "Escriba el correo institucional y la contrasena asignada.",
            "Use el boton de ojo si necesita visualizar u ocultar la contrasena.",
            "Marque la aceptacion de Habeas Data y presione INGRESAR.",
        ],
    )
    figure(
        doc,
        images["login"],
        "Figura 1. Pantalla de inicio de sesion del sistema SARC.",
        "Se resaltan el campo correo, campo contrasena, boton de visualizacion, aceptacion de datos y boton INGRESAR. La flecha indica el flujo de acceso.",
    )
    add_page_number_break(doc)

    doc.add_heading("4. Inicio y asistente virtual", level=1)
    paragraph(
        doc,
        "Al ingresar, el estudiante visualiza su nombre, cursos sugeridos y mensajes del asistente virtual. Las respuestas se muestran de forma breve y se relacionan con cursos, progreso e intereses registrados para el usuario autenticado.",
    )
    figure(
        doc,
        images["inicio"],
        "Figura 2. Pantalla de inicio con cursos sugeridos y asistente virtual.",
        "Se resaltan las recomendaciones principales y el bloque del asistente virtual, que presenta mensajes separados por perfil, avance y sugerencia.",
    )

    doc.add_heading("5. Modulo de recomendaciones", level=1)
    numbered(
        doc,
        [
            "Seleccione Recomendaciones en el menu superior.",
            "Use los filtros de area y modalidad si desea ordenar la busqueda.",
            "Revise la informacion del curso: nombre, modalidad y descripcion.",
            "Presione Ver detalles para consultar mas informacion o INSCRIBIRSE para registrar el curso.",
        ],
    )
    figure(
        doc,
        images["recomendaciones"],
        "Figura 3. Modulo de recomendaciones de cursos.",
        "Se resaltan los filtros, el area de informacion del curso y los botones principales. El fondo azul conserva la identidad institucional de SARC.",
    )
    add_page_number_break(doc)

    doc.add_heading("6. Progreso academico", level=1)
    paragraph(
        doc,
        "La seccion Mi progreso muestra el avance general del estudiante por curso. El porcentaje representa actividades completadas, recursos consultados y seguimiento academico almacenado por el sistema. Para revisar informacion adicional, use Ver reporte detallado y descargue el PDF si es necesario.",
    )
    figure(
        doc,
        images["progreso"],
        "Figura 4. Pantalla de progreso academico.",
        "Se resaltan la barra de avance, los indicadores simples y el boton para abrir el reporte PDF generado por el sistema.",
    )

    doc.add_heading("7. Perfil del estudiante", level=1)
    paragraph(
        doc,
        "El perfil permite consultar la informacion personal y academica asociada al estudiante. Desde esta pantalla se puede editar la informacion basica permitida, reiniciar los datos demo cuando aplique o cerrar la sesion.",
    )
    figure(
        doc,
        images["perfil"],
        "Figura 5. Perfil del estudiante en SARC.",
        "Se resaltan los datos del estudiante, el avatar y las acciones disponibles de la cuenta.",
    )
    add_page_number_break(doc)

    doc.add_heading("8. Recomendaciones de uso", level=1)
    bullets(
        doc,
        [
            "Revise periodicamente los cursos sugeridos para encontrar nuevas opciones academicas.",
            "Evite inscribirse en cursos que no correspondan a sus intereses o avance actual.",
            "Consulte Mi progreso para verificar actividades completadas y estado de cada curso.",
            "Descargue el reporte PDF cuando necesite presentar un resumen academico.",
            "Cierre sesion al terminar, especialmente si usa un equipo compartido.",
        ],
    )
    callout(
        doc,
        "Nota final",
        "SARC no funciona como un sistema de notas manuales. El avance mostrado se obtiene del seguimiento academico del estudiante y de la informacion registrada en la base de datos.",
        LIGHT_BLUE,
    )

    doc.core_properties.title = "Manual de Usuario SARC"
    doc.core_properties.subject = "Sistema de Apoyo Academico con Recomendacion de Cursos"
    doc.core_properties.author = "SARC"
    doc.save(OUT)


def main() -> None:
    SHOT_DIR.mkdir(parents=True, exist_ok=True)
    images = build_annotations()
    build_doc(images)
    print(OUT)


if __name__ == "__main__":
    main()
