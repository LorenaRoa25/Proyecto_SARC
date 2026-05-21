"""Genera el Manual Tecnico corto del proyecto SARC."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parents[1]
OUT = PROJECT_DIR / "Documentos" / "Manual_Tecnico_SARC.docx"

BLUE = "0B4F8A"
GREEN = "6AA84F"
TEXT = RGBColor(31, 41, 55)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = TEXT
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    for i, value in enumerate(headers):
        set_cell_text(header.cells[i], value, bold=True, size=8)
        set_cell_shading(header.cells[i], BLUE)
        header.cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            header.cells[i].width = Inches(widths[i])
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            set_cell_text(row.cells[i], value, size=8)
            if widths:
                row.cells[i].width = Inches(widths[i])
    doc.add_paragraph()


def add_page_number(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Pagina ")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def configure_doc(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(3)
        section.left_margin = Cm(4)
        section.right_margin = Cm(2)
        add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(4)

    for name in ["Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Heading 2"].font.size = Pt(11)


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.color.rgb = TEXT


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = TEXT


def cover(doc: Document) -> None:
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MANUAL TECNICO DEL SISTEMA DE APOYO ACADEMICO CON RECOMENDACION DE CURSOS (SARC)")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    doc.add_paragraph()
    for line in [
        "Documento tecnico para desarrolladores, administradores tecnicos y evaluadores academicos",
        "Lorena Roa Rivera",
        "Ingenieria Informatica",
        "Corporacion Universitaria de Sabaneta - Unisabaneta",
        "Sabaneta, Antioquia",
        "2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.name = "Arial"
        r.font.size = Pt(12)
    doc.add_page_break()


def build_doc() -> Document:
    doc = Document()
    configure_doc(doc)
    cover(doc)

    heading(doc, "1. Introduccion tecnica del sistema")
    para(doc, "El Sistema de Apoyo Academico con Recomendacion de Cursos (SARC) es una aplicacion web academica orientada a estudiantes universitarios de Unisabaneta. Su objetivo tecnico es centralizar cursos complementarios, permitir la consulta de recomendaciones, gestionar inscripciones y mostrar seguimiento de progreso desde una interfaz web modular.")
    para(doc, "El sistema resuelve la dispersion de informacion sobre cursos de apoyo y refuerzo academico. Desde el punto de vista academico, busca fortalecer competencias transversales como matematicas, programacion, comunicacion e idiomas. La version actual es funcional como prototipo web con persistencia Firebase; no implementa inteligencia artificial avanzada ni chatbot conversacional.")

    heading(doc, "2. Arquitectura del software")
    para(doc, "SARC funciona como una aplicacion web de una sola pagina con navegacion por hash. El archivo index.html contiene el contenedor principal y carga shared/js/app.js como modulo de entrada. app.js inicializa el estado, valida la sesion, protege rutas privadas y renderiza la vista correspondiente.")
    para(doc, "Aunque el proyecto no usa un framework MVC formal, se documenta como arquitectura modular inspirada en MVC: el modelo corresponde a state.db, defaultDatabase y Firestore; la vista corresponde a las plantillas HTML y hojas CSS; y los controladores son los modulos JavaScript que enlazan eventos, validan datos y ejecutan acciones.")
    para(doc, "La separacion frontend/backend se implementa mediante un frontend en HTML, CSS y JavaScript ejecutado en el navegador, y servicios administrados de Firebase para autenticacion y base de datos. app.py solo sirve archivos estaticos durante desarrollo; no es backend de negocio.")

    heading(doc, "3. Tecnologias utilizadas")
    add_table(doc, ["Tecnologia", "Funcion dentro del sistema"], [
        ["HTML5", "Estructura de index.html y plantillas de cada vista en pages/."],
        ["CSS3", "Estilos globales, layout, componentes y estilos especificos por pantalla."],
        ["JavaScript ES Modules", "Estado, rutas, eventos, reglas de negocio, recomendaciones, progreso y reportes."],
        ["Firebase Authentication", "Inicio de sesion, recuperacion y cambio de contrasena."],
        ["Cloud Firestore", "Persistencia de usuarios, cursos, recomendaciones, tareas y aceptacion de Habeas Data por userId."],
        ["Firebase SDK", "Importado dinamicamente desde gstatic en firebase-service.js."],
        ["Python http.server", "Servidor local app.py para desarrollo en localhost:8000."],
        ["localStorage", "No se usa como persistencia activa en la version revisada; la persistencia real esta en Firestore."],
    ], [1.7, 4.9])
    para(doc, "La version actual no utiliza frameworks frontend como React, Angular o Vue, ni un backend de negocio propio. La comunicacion con servicios externos se concentra en firebase-service.js para evitar repetir llamadas de autenticacion y persistencia en cada vista.")

    heading(doc, "4. Estructura de carpetas del proyecto")
    para(doc, "El proyecto esta organizado por responsabilidades. La raiz contiene index.html, app.py, README.md, FIREBASE_SETUP.md y firestore.rules. La carpeta shared agrupa estilos, modulos JavaScript y componentes reutilizables. La carpeta pages contiene las vistas funcionales del sistema. Documentos contiene entregables academicos y tools contiene scripts de apoyo documental.")
    add_table(doc, ["Ruta", "Descripcion"], [
        ["shared/css", "base.css, layout.css y components.css para estilos globales."],
        ["shared/js", "app.js, data.js, firebase-service.js, courses.js, navigation.js, pdf.js, template-loader.js y sound.js."],
        ["shared/components", "feedback.js para modales, toasts y enlaces visuales del prototipo."],
        ["pages/login", "Vista de autenticacion."],
        ["pages/recuperar", "Vista de recuperacion de contrasena."],
        ["pages/inicio", "Dashboard con cursos sugeridos y asistente."],
        ["pages/recomendaciones", "Catalogo filtrable e inscripcion."],
        ["pages/detalle, pages/progreso, pages/perfil", "Detalle de curso, seguimiento academico y gestion de perfil."],
    ], [2.0, 4.6])

    heading(doc, "5. Explicacion de archivos principales")
    add_table(doc, ["Archivo", "Responsabilidad tecnica"], [
        ["index.html", "Shell principal de la aplicacion. Carga CSS, define contenedores y llama shared/js/app.js."],
        ["shared/js/app.js", "Punto de entrada. Inicializa estado, rutas, eventos globales y renderizado de vistas."],
        ["shared/js/data.js", "Estado global, datos base, calculo de notas, recomendaciones y funciones de persistencia."],
        ["shared/js/firebase-service.js", "Conexion con Firebase Auth y Firestore; carga, guarda y resetea datos por usuario."],
        ["shared/js/firebase-config.js", "Configuracion del proyecto Firebase."],
        ["shared/js/user-profiles.js", "Centraliza reglas de perfiles especiales; define el correo demo autorizado y valida funciones exclusivas."],
        ["shared/js/courses.js", "Filtros de cursos, busqueda por id, inscripcion, cupos y refresco del dashboard."],
        ["shared/js/pdf.js", "Genera reportes PDF simples en el navegador sin libreria externa."],
        ["firestore.rules", "Reglas de seguridad para aislar documentos por userId autenticado."],
        ["app.py", "Servidor local de archivos estaticos con MIME types y cabeceras de desarrollo."],
    ], [2.0, 4.6])

    heading(doc, "6. Flujo del sistema")
    para(doc, "El flujo inicia cuando el navegador carga index.html y app.js inicializa Firebase. Si el usuario no esta autenticado, se renderiza login o recuperacion. Si la sesion existe, se habilita el dashboard y las vistas internas.")
    bullets(doc, [
        "Autenticacion: login.js valida formulario, usa loginWithEmail() y carga datos con loadDatabaseForUser().",
        "Recomendaciones: inicio.js y data.js generan sugerencias con base en cursos inscritos, actividades y rendimiento.",
        "Inscripcion: courses.js valida curso, cupos y estado registered; luego actualiza state.db y guarda en Firestore.",
        "Perfil demo: perfil.js muestra el boton Reiniciar demo solo si la sesion y el perfil corresponden al correo autorizado.",
        "Navegacion: navigation.js modifica window.location.hash y app.js renderiza la vista correspondiente.",
    ])

    heading(doc, "7. Gestion de recomendaciones de cursos")
    para(doc, "La recomendacion de cursos en SARC es una funcionalidad implementada con reglas de negocio y datos academicos disponibles. No corresponde a machine learning. Los cursos base se filtran mediante CORE_COURSE_IDS y la vista Recomendaciones permite filtrar por area interdisciplinaria y modalidad.")
    para(doc, "El asistente virtual es prototipado: presenta recomendaciones textuales segun rendimiento calculado, cursos inscritos y estado academico. No responde preguntas en lenguaje natural ni se conecta a servicios externos de IA.")

    heading(doc, "8. Modelo de datos")
    add_table(doc, ["Entidad", "Representacion real en el proyecto"], [
        ["Usuario", "Documento usuarios/{uid}; incluye email, name, faculty, career, semester, modality, avatar y userId."],
        ["Curso", "Documento cursos/{uid}_{courseId}; incluye area, modalidad, cupos, progreso, actividades y estado."],
        ["Recomendacion", "Documento recomendaciones/{uid}_{courseId}; guarda courseId, courseName, recommendation y alternatives."],
        ["Inscripcion", "No existe coleccion separada; se representa en cada curso con enrolled y registered."],
        ["Progreso", "No existe coleccion separada; se almacena dentro del curso con progress, lessons, activities y average."],
    ], [1.4, 5.2])

    heading(doc, "9. Ficha tecnica de hardware")
    add_table(doc, ["Elemento", "Requisito minimo", "Requisito recomendado"], [
        ["Procesador cliente", "Dual Core 1.8 GHz para navegacion basica.", "Intel i3/Ryzen 3 o superior para uso fluido y pruebas."],
        ["Memoria RAM cliente", "4 GB.", "8 GB o superior."],
        ["Almacenamiento cliente", "100 MB libres para cache, documentos y descargas PDF.", "500 MB o mas para evidencias, capturas y documentacion."],
        ["Equipo de desarrollo", "Windows, macOS o Linux con Python 3 y navegador moderno.", "Equipo con 8 GB RAM, editor de codigo y conexion estable."],
        ["Servidor / hosting", "No requiere servidor propio de negocio; puede publicarse como sitio estatico.", "Hosting HTTPS o Firebase Hosting con Firebase Authentication y Firestore configurados."],
        ["Conectividad", "Internet estable para Firebase.", "Conexion institucional estable para pruebas y sustentacion."],
    ], [1.55, 2.55, 2.55])

    heading(doc, "10. Ficha tecnica de software")
    add_table(doc, ["Categoria", "Detalle"], [
        ["Tecnologias", "HTML5, CSS3, JavaScript ES Modules, Firebase Authentication, Cloud Firestore y Python http.server."],
        ["Versiones", "Firebase SDK 10.12.5 importado dinamicamente; Python 3 para app.py; navegadores modernos compatibles con ES Modules."],
        ["Firebase", "Autenticacion Email/Password, recuperacion y cambio de contrasena; Firestore para usuarios, cursos, recomendaciones, tareas y Habeas Data."],
        ["Librerias", "No se requieren dependencias npm obligatorias ni frameworks frontend; el SDK de Firebase se consume desde CDN."],
        ["Compatibilidad", "Chrome, Edge, Firefox y Safari modernos con JavaScript habilitado, fetch, Blob y FileReader."],
        ["Hosting", "Ejecucion local con iniciar_sarc.bat y app.py; despliegue recomendado en hosting estatico HTTPS o Firebase Hosting."],
        ["Dependencias del proyecto", "shared/js/firebase-config.js, firestore.rules y conexion a internet para servicios Firebase."],
    ], [1.7, 4.9])

    heading(doc, "11. Instalacion y despliegue")
    bullets(doc, [
        "Clonar o copiar el proyecto completo conservando la estructura de carpetas.",
        "Configurar Firebase en shared/js/firebase-config.js con los datos del proyecto institucional.",
        "Activar Firebase Authentication con Email/Password y crear Cloud Firestore.",
        "Publicar firestore.rules para proteger datos por userId.",
        "Ejecutar localmente con iniciar_sarc.bat y abrir http://localhost:8000/.",
        "Para despliegue basico, publicar los archivos estaticos en un hosting compatible con HTTPS, por ejemplo Firebase Hosting.",
    ])

    heading(doc, "12. Seguridad del sistema")
    para(doc, "El sistema valida formularios de autenticacion, contrasena y perfil desde el frontend. Las credenciales se administran mediante Firebase Authentication y no se almacenan en Firestore. La aceptacion de Habeas Data se registra en la coleccion habeasData.")
    para(doc, "Las reglas de Firestore verifican que el usuario autenticado coincida con el userId del documento consultado o modificado. Esto protege usuarios, cursos, recomendaciones y tareasAsistente contra acceso cruzado. Como mejora futura se recomienda auditoria de reglas, monitoreo y manejo institucional de roles.")
    para(doc, "La funcionalidad Reiniciar demo se protege por perfil: el boton se renderiza unicamente para lorena.roa@unisabaneta.edu.co y resetSarcDemo() valida nuevamente el correo autenticado antes de ejecutar el reinicio. Esto evita que Fabiana u otros usuarios nuevos accedan a la herramienta de demostracion.")

    heading(doc, "13. Consideraciones de mantenimiento")
    para(doc, "La modularidad del proyecto facilita mantenimiento porque cada vista posee sus archivos HTML, CSS y JS, mientras la logica compartida se concentra en shared/js. Para evolucionar el sistema se recomienda conservar esta separacion, documentar nuevos modulos y evitar mezclar logica de Firebase directamente en las vistas.")
    bullets(doc, [
        "Mejoras futuras: registro formal, panel administrativo, Firebase Storage para avatares y pruebas automatizadas.",
        "Escalabilidad: Firestore permite crecer por documentos asociados a userId; se recomienda revisar indices y reglas si aumenta el volumen.",
        "Calidad: agregar pruebas para login, filtros, inscripcion, progreso, perfil y reglas de seguridad.",
    ])

    heading(doc, "14. Conclusiones tecnicas")
    para(doc, "SARC es un sistema web academico funcional, modular y comprensible para un contexto universitario. Su principal valor tecnico esta en integrar una interfaz por vistas con servicios Firebase para autenticacion y persistencia, manteniendo separacion entre estado, navegacion, cursos, reportes y componentes de feedback.")
    para(doc, "El sistema aporta al acompanamiento academico mediante centralizacion de cursos, recomendaciones basadas en reglas, seguimiento de progreso y reportes descargables. Su potencial de crecimiento esta en ampliar roles, mejorar el motor de recomendacion, automatizar pruebas y formalizar un despliegue institucional con HTTPS y monitoreo.")

    core = doc.core_properties
    core.title = "Manual Tecnico SARC"
    core.subject = "Manual tecnico del Sistema de Apoyo Academico con Recomendacion de Cursos"
    core.author = "Lorena Roa Rivera"
    core.keywords = "SARC, manual tecnico, Firebase, HTML, CSS, JavaScript"
    return doc


def main() -> None:
    doc = build_doc()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
